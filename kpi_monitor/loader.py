"""Consolidate workbook rows without changing their values."""

import logging
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, List, Tuple

from openpyxl import load_workbook

LOGGER = logging.getLogger(__name__)


@dataclass
class LoadedData:
    headers: Tuple[Any, ...]
    rows: List[Tuple[Any, ...]]


def _day(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return datetime.fromisoformat(value.strip()).date()
        except ValueError:
            pass
    raise ValueError("Fecha invalida para comparar duplicados")


def load_data(config):
    """Read every sheet, in sorted file order and original sheet/row order."""
    input_dir = Path(config["input_dir"])
    input_dir.mkdir(parents=True, exist_ok=True)
    Path(config["output_dir"]).mkdir(parents=True, exist_ok=True)
    files = sorted(path for path in input_dir.glob("*.xlsx")
                   if path.is_file() and not path.name.startswith("~$"))
    if not files:
        raise ValueError("No hay archivos xlsx en " + str(input_dir))
    required = [config["site_column"], config["date_column"]] + list(config["kpis"].values())
    headers = None
    rows = []
    seen = set()
    duplicates = set()
    examples = []
    for path in files:
        with path.open("rb") as source:
            workbook = load_workbook(source, read_only=True, data_only=False)
            try:
                for sheet in workbook.worksheets:
                    iterator = sheet.iter_rows(values_only=True)
                    current = tuple(next(iterator, ()))
                    for column in required:
                        if column not in current:
                            raise ValueError("Falta columna {!r} en {} ({})".format(column, path, sheet.title))
                    if len(set(current)) != len(current):
                        raise ValueError("Encabezados repetidos en {} ({})".format(path, sheet.title))
                    if headers is None:
                        headers = current
                    if set(current) != set(headers):
                        raise ValueError("Encabezados incompatibles en {} ({})".format(path, sheet.title))
                    indices = [current.index(column) for column in headers]
                    site_index = headers.index(config["site_column"])
                    date_index = headers.index(config["date_column"])
                    for number, row in enumerate(iterator, start=2):
                        values = tuple(row[index] for index in indices)
                        try:
                            key = (str(values[site_index]).strip(), _day(values[date_index]))
                        except ValueError as error:
                            raise ValueError("Fecha invalida en {} ({}, fila {})".format(path, sheet.title, number)) from error
                        if key in seen and key not in duplicates:
                            duplicates.add(key)
                            if len(examples) < 5:
                                examples.append(key)
                        seen.add(key)
                        rows.append(values)
            finally:
                workbook.close()
    if duplicates:
        LOGGER.warning("%d claves sitio+fecha duplicadas; ejemplos: %s", len(duplicates), examples)
    return LoadedData(headers, rows)
